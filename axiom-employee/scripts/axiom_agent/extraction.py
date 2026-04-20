"""Bounded, local parsers. Macros and formulas are never executed."""
import csv
import io
import zipfile
import uuid
from xml.parsers import expat
from pathlib import Path
from .models import Document, Section

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_SECTIONS = 50_000

class ExtractionError(ValueError): pass

def validate_office_xml(data):
    # Streaming preflight rejects DTDs/entities in any encoding before library parsing.
    parser = expat.ParserCreate()
    depth, nodes = 0, 0
    def forbidden(*args): raise ExtractionError("Office XML declarations and entities are unsupported")
    def start(*args):
        nonlocal depth, nodes
        depth += 1
        nodes += 1
        if depth > 256 or nodes > 500000: raise ExtractionError("Office XML complexity exceeds limits")
    def end(*args):
        nonlocal depth
        depth -= 1
    parser.StartDoctypeDeclHandler = forbidden
    parser.EntityDeclHandler = forbidden
    parser.ExternalEntityRefHandler = forbidden
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.Parse(data, True)


def extract(name: str, data: bytes) -> Document:
    kind = Path(name).suffix.lower().lstrip(".")
    if kind not in {"txt", "csv", "pdf", "docx", "xlsx"}:
        raise ExtractionError("Supported formats: PDF, DOCX, XLSX, TXT, CSV")
    if not data or len(data) > MAX_BYTES:
        raise ExtractionError("File is empty or exceeds 20 MiB")
    sections = []
    total = 0
    def add(text, location):
        nonlocal total
        total += len(text)
        if total > MAX_TEXT or len(sections) >= MAX_SECTIONS:
            raise ExtractionError("Document exceeds extraction limits")
        if text: sections.append(Section(text, location))
    try:
        if kind in {"docx", "xlsx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                    raise ExtractionError("Office archive exceeds decompression limits")
                if any("vbaproject" in e.filename.lower() for e in entries):
                    raise ExtractionError("Macro-bearing documents are unsupported")
                for entry in entries:
                    if entry.filename.lower().endswith(".xml"):
                        validate_office_xml(archive.read(entry))
        if kind == "txt":
            add(data.decode("utf-8-sig"), {"offset": 0})
        elif kind == "csv":
            for row, cells in enumerate(csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True), 1):
                for col, text in enumerate(cells, 1): add(text, {"row": row, "column": col})
        elif kind == "pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.needs_pass: raise ExtractionError("Encrypted PDFs are unsupported")
                if len(pdf) > 1000: raise ExtractionError("PDF exceeds page limit")
                for index, page in enumerate(pdf):
                    add(page.get_text("text"), {"page": index + 1})
        elif kind == "docx":
            from docx import Document as WordDocument
            doc = WordDocument(io.BytesIO(data))
            # Paragraph offsets remain reliable even when a match crosses runs.
            for index, para in enumerate(doc.paragraphs): add(para.text, {"paragraph": index + 1})
            for ti, table in enumerate(doc.tables, 1):
                for ri, row in enumerate(table.rows, 1):
                    for ci, cell in enumerate(row.cells, 1):
                        add(cell.text, {"table": ti, "row": ri, "column": ci})
            for si, section in enumerate(doc.sections, 1):
                for part, container in (("header", section.header), ("footer", section.footer)):
                    for pi, para in enumerate(container.paragraphs, 1):
                        add(para.text, {"word_section": si, "part": part, "paragraph": pi})
        elif kind == "xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in book:
                    if sheet.max_row > MAX_SECTIONS or sheet.max_column > 1000:
                        raise ExtractionError("Worksheet dimensions exceed limits")
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value is not None:
                                add(str(cell.value), {"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.data_type == "f"})
            finally: book.close()
    except ExtractionError: raise
    except ImportError:
        raise ExtractionError(f"Local {kind.upper()} parser is unavailable in Axiom's Python environment. Install scripts/requirements.txt using the interpreter configured by AXIOM_PYTHON, then restart Axiom.") from None
    except Exception:
        raise ExtractionError("Document is malformed or unreadable") from None
    if not sections:
        raise ExtractionError("No extractable text; image-only documents need local OCR (not implemented)")
    return Document(str(uuid.uuid4()), Path(name.replace("\\", "/")).name, kind, sections,
                    {"coverage": "Extracted text only; embedded media, comments, annotations and hidden package data are not covered."})
