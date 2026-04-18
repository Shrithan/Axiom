import base64
import io
import json
import socket
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from axiom_agent.extraction import extract, ExtractionError
from axiom_agent.models import ContentEvent
from axiom_agent.detectors import PrivacyDetector, SecretDetector, OrganizationDetector
from axiom_agent.policy import default_rules, validate_rules
from axiom_agent.pipeline import Pipeline
from axiom_agent.service import Service
from axiom_agent.storage import LocalStore


def custom(**changes):
    rule = dict(id='project', name='Project Penguin', kind='keyword', pattern='penguin',
                case_sensitive=False, severity='HIGH', action='REDACT', transformation='REDACT',
                replacement='[REMOVED]', enabled=True, destinations=[], events=[])
    rule.update(changes)
    return rule


def scan(text, rules=None, context=None):
    return Pipeline(default_rules() + [custom()] if rules is None else rules).scan('sample.txt', text.encode(), context)


class DetectorTests(unittest.TestCase):
    def privacy(self, text): return PrivacyDetector().detect(extract('a.txt', text.encode()))
    def test_email_phone_ip(self):
        kinds = {d.type for d in self.privacy('Contact john@example.com +1 (312) 555-1234 at 192.168.1.4')}
        self.assertTrue({'EMAIL', 'PHONE', 'IP_ADDRESS'} <= kinds)
    def test_ssn_validation(self):
        self.assertEqual([d.matched_text for d in self.privacy('000-12-1234 666-12-1234 900-12-1234 123-00-1234 123-12-0000 123-45-6789') if d.type == 'SSN'], ['123-45-6789'])
    def test_luhn_and_false_positives(self):
        matches = [d.matched_text for d in self.privacy('4111 1111 1111 1111; 4111 1111 1111 1112; 0000000000000000; 999.1.1.1; 111-111-1111')]
        self.assertEqual(matches, ['4111 1111 1111 1111'])
    def test_keyword_case(self):
        self.assertEqual(len(scan('Penguin penguin PENGUIN', [custom()])['detections']), 3)
        self.assertEqual(len(scan('Penguin penguin', [custom(case_sensitive=True)])['detections']), 1)
    def test_unicode_offsets(self):
        r = scan('🔒 Project Penguin', [custom()])
        self.assertEqual(r['protected_sections'][0]['text'], '🔒 Project ███████')
    def test_regex_rule(self):
        r = scan('Customer ACME-123', [custom(kind='regex', pattern=r'ACME-\d{3}')])
        self.assertEqual(r['detections'][0]['matched_text'], 'ACME-123')
    def test_secret_patterns(self):
        text = 'AKIA' + 'A'*16 + ' ghp_' + 'a'*36 + '\nBearer ' + 'b'*24 + '\npostgres://user:secret@db\npassword=longsecret\n-----BEGIN PRIVATE KEY-----\nabc\n-----END PRIVATE KEY-----'
        kinds = {d.type for d in SecretDetector().detect(extract('a.txt', text.encode()))}
        self.assertEqual(kinds, {'API_KEY','BEARER_TOKEN','URL_CREDENTIALS','CONNECTION_STRING','PRIVATE_KEY'})
    def test_bad_regex_and_timeout(self):
        with self.assertRaises(ValueError): validate_rules([custom(kind='regex', pattern='[')])
        with self.assertRaises(ValueError): validate_rules([custom(kind='regex', pattern='a*')])
        with self.assertRaisesRegex(ValueError, 'execution budget'):
            scan('a'*100000 + '!', [custom(kind='regex', pattern='(a+)+$')])
    def test_keyword_is_literal(self):
        self.assertEqual(len(scan('a.b axb', [custom(pattern='a.b')])['detections']),1)


class PolicyTests(unittest.TestCase):
    def test_penguin_integration(self):
        r = scan('Project Penguin launches on November 3. Contact john@example.com for details.')
        self.assertEqual(r['status'], 'REDACTED')
        self.assertEqual(r['protected_sections'][0]['text'], 'Project ███████ launches on November 3. Contact ' + '█'*16 + ' for details.')
        self.assertEqual({d['rule_name'] for d in r['decisions']}, {'Project Penguin','Email privacy'})
    def test_precedence(self):
        rules = [custom(id='allow', action='ALLOW'), custom(id='warn',action='WARN'), custom(id='redact'), custom(id='block',action='BLOCK')]
        r = scan('penguin', rules)
        self.assertEqual(r['action'], 'BLOCK'); self.assertIsNone(r['protected_sections'])
    def test_builtin_competing_policies(self):
        rules = [custom(id='a',kind='builtin',pattern='EMAIL',action='ALLOW'), custom(id='b',kind='builtin',pattern='EMAIL',action='BLOCK')]
        self.assertEqual(scan('john@example.com',rules)['decisions'][0]['rule_id'],'b')
    def test_tie_break_is_stable(self):
        a,b = custom(id='a',transformation='REMOVE'),custom(id='b',transformation='REPLACE')
        self.assertEqual(scan('penguin',[a,b])['protected_sections'],scan('penguin',[b,a])['protected_sections'])
    def test_disabled(self):
        self.assertEqual(scan('Penguin',[custom(enabled=False)])['detections'],[])
        rules = default_rules()
        for r in rules: r['enabled']=False
        result=scan('john@example.com',rules)
        self.assertEqual(result['action'],'ALLOW');self.assertEqual(len(result['detections']),1)
    def test_context(self):
        rules=[custom(action='BLOCK',destinations=['PUBLIC'],events=['EXPORT'])]
        self.assertEqual(scan('penguin',rules,ContentEvent(destination='INTERNAL'))['action'],'ALLOW')
        self.assertEqual(scan('penguin',rules,ContentEvent(action='EXPORT',destination='PUBLIC'))['action'],'BLOCK')
    def test_allow_warn(self):
        for action in ['ALLOW','WARN']:
            r=scan('Penguin',[custom(action=action)])
            self.assertEqual(r['action'],action);self.assertEqual(r['protected_sections'][0]['text'],'Penguin')
    def test_mask(self):
        r=scan('123-45-6789',[custom(kind='builtin',pattern='SSN',transformation='MASK')])
        self.assertEqual(r['protected_sections'][0]['text'],'***-**-6789')
    def test_remove_replace_multiple(self):
        for mode,expected in [('REMOVE',' and '),('REPLACE','[REMOVED] and [REMOVED]')]:
            self.assertEqual(scan('penguin and penguin',[custom(transformation=mode)])['protected_sections'][0]['text'],expected)
    def test_overlaps_union(self):
        rules=[custom(pattern='Project Penguin',transformation='REPLACE'),custom(id='inner',transformation='MASK')]
        self.assertEqual(scan('Project Penguin!',rules)['protected_sections'][0]['text'],'███████ ███████!')
    def test_block_all_output(self):
        self.assertIsNone(scan('123-45-6789 and penguin')['protected_sections'])
    def test_invalid_context(self):
        with self.assertRaises(ValueError):scan('penguin',context=ContentEvent(destination='UNKNOWN'))


class ExtractionTests(unittest.TestCase):
    def test_txt_csv_locations(self):
        self.assertEqual(extract('a.txt',b'hello').sections[0].location,{'offset':0})
        d=extract('a.csv',b'name,project\nJohn,"Penguin, confidential"\n')
        self.assertEqual(d.sections[-1].location,{'row':2,'column':2})
        self.assertEqual(d.sections[-1].text,'Penguin, confidential')
    def test_docx_runs_table_header(self):
        from docx import Document
        d=Document();p=d.add_paragraph('Project ');p.add_run('Peng');p.add_run('uin')
        d.add_table(rows=1,cols=1).cell(0,0).text='john@example.com'
        d.sections[0].header.paragraphs[0].text='Penguin header'
        raw=io.BytesIO();d.save(raw)
        doc=extract('a.docx',raw.getvalue())
        result=Pipeline(default_rules()+[custom()]).process(doc)
        self.assertEqual(len(result['detections']),3)
        self.assertEqual(doc.sections[0].location,{'paragraph':1})
    def test_xlsx_sheet_cell_formula(self):
        from openpyxl import Workbook
        b=Workbook();s=b.active;s.title='Projects';s['C2']='Project Penguin';s['D2']='=SUM(1,2)'
        raw=io.BytesIO();b.save(raw)
        doc=extract('a.xlsx',raw.getvalue())
        self.assertEqual(doc.sections[0].location['cell'],'C2')
        self.assertEqual(doc.sections[0].location['sheet'],'Projects')
        self.assertTrue(doc.sections[1].location['formula'])
        self.assertEqual(Pipeline([custom()]).process(doc)['action'],'REDACT')
    def test_pdf_page(self):
        import fitz
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72,72),'Project Penguin john@example.com')
            raw=pdf.tobytes()
        result=Pipeline(default_rules()+[custom()]).scan('a.pdf',raw)
        self.assertEqual(len(result['detections']),2)
        self.assertEqual(result['detections'][0]['location']['page'],1)
    def test_malformed_and_unsupported(self):
        for name,data in [('a.pdf',b'not pdf'),('a.docx',b'bad'),('a.xlsx',b'bad'),('a.txt',b'\xff'),('a.csv',b'"unterminated'),('a.doc',b'a'),('a.txt',b'')]:
            with self.subTest(name=name),self.assertRaises(ExtractionError):extract(name,data)
    def test_office_xml_entities_are_rejected(self):
        for data in [b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>',
                     '<!DOCTYPE x [<!ENTITY secret "value">]><x>&secret;</x>'.encode('utf-16')]:
            raw=io.BytesIO()
            with zipfile.ZipFile(raw,'w') as archive:archive.writestr('word/document.xml',data)
            with self.assertRaisesRegex(ExtractionError,'declarations and entities'):
                extract('a.docx',raw.getvalue())
    def test_image_only_pdf(self):
        import fitz
        with fitz.open() as pdf:pdf.new_page();raw=pdf.tobytes()
        with self.assertRaisesRegex(ExtractionError,'OCR'):extract('a.pdf',raw)
    def test_archive_limits_and_macros(self):
        for filename,data in [('word/vbaProject.bin',b'code'),('large',b'a'*(81*1024*1024))]:
            raw=io.BytesIO()
            with zipfile.ZipFile(raw,'w',zipfile.ZIP_DEFLATED) as z:z.writestr(filename,data)
            with self.assertRaises(ExtractionError):extract('a.docx',raw.getvalue())
    def test_sanitized_name(self):
        self.assertEqual(extract('../../private.txt',b'hello').name,'private.txt')


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.store=LocalStore(self.temp.name);self.service=Service(self.store)
    def tearDown(self):self.temp.cleanup()
    def req(self,text,op='scan',**kwargs):
        return self.service.handle(dict(op=op,name='private.txt',data=base64.b64encode(text.encode()).decode(),**kwargs))
    def test_persist_and_audit_minimization(self):
        self.store.save_rules(default_rules()+[custom()]);self.req('Penguin john@example.com')
        self.assertEqual(LocalStore(self.temp.name).rules()[-1]['name'],'Project Penguin')
        audit=(Path(self.temp.name)/'audit.jsonl').read_text()
        for raw in ['Penguin','john@example.com','private.txt','matched_text']:self.assertNotIn(raw,audit)
        self.assertEqual(len(self.store.events()),2)
    def test_export_rechecks_policy_and_event(self):
        self.store.save_rules([custom(action='BLOCK',events=['EXPORT'])])
        self.assertEqual(self.req('Penguin')['action'],'ALLOW')
        with self.assertRaisesRegex(ValueError,'blocked'):self.req('Penguin','export')
    def test_export_warning_and_block_name(self):
        self.store.save_rules([custom(action='WARN',events=['EXPORT'])])
        result=self.req('Penguin','export')
        self.assertIn('Policy warning: Project Penguin',result['warnings'])
        self.store.save_rules([custom(action='BLOCK',events=['EXPORT'])])
        with self.assertRaisesRegex(ValueError,'Project Penguin'):self.req('Penguin','export')
    def test_export_rechecks_changed_rules(self):
        self.store.save_rules([custom()]);self.assertEqual(self.req('Penguin')['action'],'REDACT')
        self.store.save_rules([custom(action='BLOCK')])
        with self.assertRaises(ValueError):self.req('Penguin','export')
    def test_export_transforms(self):
        self.store.save_rules(default_rules()+[custom()]);out=self.req('Penguin john@example.com','export')
        self.assertEqual(out['extension'],'txt');self.assertNotIn('Penguin',out['text']);self.assertNotIn('john@',out['text'])
    def test_csv_export_shape_and_injection(self):
        self.store.save_rules([custom()])
        out=self.service.handle(dict(op='export',name='a.csv',data=base64.b64encode(b'a,,c\nPenguin,=SUM(1),\n').decode()))
        self.assertEqual(out['extension'],'csv');self.assertIn("'",out['text']);self.assertIn('a,,c',out['text'])
    def test_no_network(self):
        with patch.object(socket,'socket',side_effect=AssertionError('Network forbidden')):
            self.store.save_rules([custom()]);self.req('Project Penguin');self.req('Project Penguin','export')
    def test_invalid_base64(self):
        with self.assertRaises(ValueError):self.service.handle(dict(op='scan',name='a.txt',data='***'))


class BoundaryTests(unittest.TestCase):
    def handler(self, headers=None, payload=b''):
        from local_server import Handler
        handler = object.__new__(Handler)
        handler.path='/api/agent'
        handler.headers=headers or {}
        handler.rfile=io.BytesIO(payload)
        handler.response=None
        handler.respond=lambda status,body:setattr(handler,'response',(status,body))
        return handler
    def test_loopback_host_token_and_origin(self):
        from local_server import TOKEN, ORIGIN
        correct={'Host':'127.0.0.1:8765','Origin':ORIGIN,'X-Axiom-Token':TOKEN,'Content-Length':'2'}
        for field in ['Host','Origin','X-Axiom-Token']:
            headers={**correct,field:'untrusted'}
            h=self.handler(headers,b'{}');h.do_POST();self.assertEqual(h.response[0],403)
    def test_http_request_size(self):
        from local_server import TOKEN,ORIGIN
        h=self.handler({'Host':'127.0.0.1:8765','Origin':ORIGIN,'X-Axiom-Token':TOKEN,'Content-Length':str(40*1024*1024)})
        h.do_POST();self.assertEqual(h.response[0],400)
    def test_http_worker_result(self):
        import subprocess
        from local_server import TOKEN,ORIGIN
        raw=b'{"op":"rules"}'
        h=self.handler({'Host':'127.0.0.1:8765','Origin':ORIGIN,'X-Axiom-Token':TOKEN,'Content-Length':str(len(raw))},raw)
        with patch('local_server.subprocess.run',return_value=subprocess.CompletedProcess([],0,b'{"result":[]}',b'')) as worker:
            h.do_POST();self.assertEqual(h.response,(200,{'result':[]}));self.assertEqual(worker.call_args.kwargs['timeout'],30)
    def test_ipc_and_disk_export_artifact(self):
        import os
        import subprocess
        with tempfile.TemporaryDirectory() as temp:
            LocalStore(temp).save_rules(default_rules()+[custom()])
            req=dict(op='export',name='penguin.txt',data=base64.b64encode(b'Project Penguin john@example.com').decode())
            completed=subprocess.run([sys.executable,str(Path(__file__).resolve().parents[1]/'scripts/agent_ipc.py')],input=json.dumps(req),text=True,capture_output=True,env={**os.environ,'AXIOM_DATA_DIR':temp},timeout=30)
            self.assertEqual(completed.returncode,0)
            output=json.loads(completed.stdout)['result']['text']
            path=Path(temp)/'penguin.protected.txt';path.write_text(output)
            self.assertEqual(path.read_text(),'Project ███████ '+'█'*16)
            self.assertEqual(completed.stderr,'')
    def test_all_formats_penguin_pipeline(self):
        import fitz
        from docx import Document
        from openpyxl import Workbook
        text='Project Penguin launches Friday. Contact john@example.com.'
        files=[('a.txt',text.encode()),('a.csv',('content\n'+text+'\n').encode())]
        with fitz.open() as pdf:pdf.new_page().insert_text((20,72),text);files.append(('a.pdf',pdf.tobytes()))
        doc=Document();doc.add_paragraph(text);raw=io.BytesIO();doc.save(raw);files.append(('a.docx',raw.getvalue()))
        book=Workbook();book.active['A1']=text;raw=io.BytesIO();book.save(raw);files.append(('a.xlsx',raw.getvalue()))
        for name,data in files:
            with self.subTest(name=name):
                result=Pipeline(default_rules()+[custom()]).scan(name,data)
                self.assertEqual(len(result['detections']),2)
                protected=' '.join(s['text'] for s in result['protected_sections'])
                self.assertNotIn('Penguin',protected);self.assertNotIn('john@example.com',protected)

if __name__=='__main__':unittest.main()
